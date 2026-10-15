"""Phase-1 rebuild of NeuroImaging_CAFN_Presentation.pptx.

Reads  docs/NeuroImaging_CAFN_Presentation_v1_backup.pptx  (never modified)
Writes docs/NeuroImaging_CAFN_Presentation.pptx

Run:  .venv/bin/python docs/pptx_tools/build_phase1.py
"""
import os, sys
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Emu
from pptx_lib import *

DOCS = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
SRC = os.path.join(DOCS, 'NeuroImaging_CAFN_Presentation_v1_backup.pptx')
DST = os.path.join(DOCS, 'NeuroImaging_CAFN_Presentation.pptx')
ARCH = os.path.join(DOCS, 'journal_ieee_access', 'figs', 'fig_arch_v2.png')

prs = Presentation(SRC)
S = list(prs.slides)          # S[0] = original slide 1 ... S[20] = original slide 21
fac = SlideFactory(prs, S[14])

# icons harvested from the original deck
ICON = {}
for s in S:
    for sh in s.shapes:
        if sh.shape_type == 13 and sh.name.endswith(' icon'):
            ICON[sh.name[:-5]] = sh.image.blob


def txt(s, prefix, nth=0):
    return find(s, prefix=prefix, nth=nth)


def banner(s, name, x, y, w, h, paras, fill='tx2', adj=10526, pad=365760):
    shape(s, name, x, y, w, h, 'roundRect', adj, fill)
    textbox(s, name + ' text', x + pad, y + 45720, w - 2 * pad, h - 91440, paras, anchor='ctr')


# =============================================================================
# Slide 1  cover
# =============================================================================
s = S[0]
set_text(txt(s, 'Detecting and classifying'),
         'Detecting and classifying vision defects: a controlled synthetic benchmark and reproducible multi-cohort evaluation')
set_text(txt(s, '16 temporal tokens'), '16 temporal tokens · 1D VEP / spectrum')
set_notes(s, "Good morning. This is the NeuroImaging project: CAFN, a Cross-Attention Fusion Network that combines a 3D structural MRI volume "
    "with a 1D signal, a visual evoked potential in the synthetic study and a resting-state power spectrum in the real cohorts. The picture on the right "
    "is the whole idea in one frame: the MRI becomes 64 spatial tokens, the signal becomes 16 temporal tokens, and in the current design the two token "
    "sets query each other through gated bidirectional cross-attention. The project has two parts: a controlled synthetic benchmark that asks when "
    "fusion helps, and a fully reproducible evaluation on two public real cohorts. One framing point up front: no public dataset has paired MRI, VEP and "
    "eye-disease labels, so I make no diagnostic accuracy claim, and the real-cohort results are at chance level.")

# =============================================================================
# Slide 2  motivation
# =============================================================================
s = S[1]
set_text(txt(s, 'Can each spatial region'),
         'Can each spatial region of the brain scan ask which temporal features of the signal are relevant to it, and can the signal ask which regions matter?')
set_notes(s, "Disorders of the visual pathway, such as glaucoma, optic neuritis and cortical visual impairment, can be examined in two ways. "
    "Structural MRI tells us where: atrophy, lesions, regional change, but it carries no timing information. The visual evoked potential tells us when: "
    "P100 latency and amplitude reflect conduction, but there is no spatial localisation. Glaucoma alone affected an estimated 64.3 million people in 2013. "
    "Our question is whether the two can be fused so that each spatial region asks which temporal features matter to it. In the final design the "
    "signal also queries the scan, which is why the fusion is bidirectional. The diagram on the right shows the original one-direction idea.")

# =============================================================================
# Slide 3  four obstacles
# =============================================================================
s = S[2]
set_text(txt(s, 'One network, two streams'),
         'One network, two streams: an SE 3D residual CNN for MRI and a multi-scale 1D CNN for the signal, with a fused head and two auxiliary heads.')
set_text(txt(s, 'Cross-attention: spatial'),
         'Gated bidirectional cross-attention: each stream queries the other, with learned gates that start half open.')
set_text(txt(s, '3D Grad-CAM and an attention'),
         '3D Grad-CAM and Grad-CAM++ plus per-head attention maps, used as inspection aids rather than explanations.')
set_text(txt(s, 'A class-conditioned'),
         'A synthetic generator with tunable modality information, plus two public MRI + EEG/MEG cohorts as the closest real test.')
set_notes(s, "There are four obstacles. First, most pipelines analyse either imaging or electrophysiology, not both: our answer is one network with two "
    "encoders and a fused head plus two auxiliary heads. Second, a volume and a waveform differ in dimensionality, sampling and noise, so naive concatenation "
    "gives every feature the same standing: our answer is gated bidirectional cross-attention. Third, black-box classifiers are hard to trust clinically: "
    "we provide 3D Grad-CAM, Grad-CAM++ and attention maps, strictly as inspection aids. Fourth, there is no public dataset with paired MRI, VEP and "
    "labels: we use a controlled synthetic generator where the information in each modality is a knob, and two public real cohorts as the closest available real test.")

# =============================================================================
# NEW: What is new in this project
# =============================================================================
N_new = fac.new('What is new in this project')
cw, ch = 5410048, 2011680
items = [
    ('Architecture', 'accent1', 'Cross-modal heterogeneity',
     'Gated bidirectional cross-attention couples an SE 3D residual encoder and a multi-scale 1D encoder, with learned attention pooling, positional embeddings and two auxiliary heads. Deterministic pooling makes GPU training reproducible.'),
    ('When fusion helps', 'accent3', 'Add baselines',
     'A synthetic generator with independent knobs for the information in each modality: like-for-like comparison with concatenation, sum and unimodal models, a component ablation and an informativeness sweep, over several seeds.'),
    ('Reproducible real data', 'accent4', 'Train properly',
     'One command downloads two public OpenNeuro cohorts, preprocesses, trains and tests: pooled and leave-one-cohort-out models, repeated CV, bootstrap intervals, permutation tests, classical baselines, one fixed seed.'),
    ('Candid findings', 'accent2', 'Check explanations',
     'We report where the design helps and where it does not. On the real cohorts the signal is at most weak and does not transfer between cohorts, and we list the pitfalls we found along the way.'),
]
for k, (title, col, icon, body) in enumerate(items):
    x = L if k % 2 == 0 else 6233008
    y = TOP if k < 2 else 3520440
    card(N_new, f'contribution {k+1}', x, y, cw, ch, 4000)
    badge(N_new, title, x + 274320, y + 228600, 548640, col, ICON[icon])
    textbox(N_new, f'contribution {k+1} title', x + 1005840, y + 228600, cw - 1280160, 548640,
            [para([run(f'{k+1}  {title}', 20, b=True, color='tx1', font='Cambria')], sz_end=20)], anchor='ctr')
    textbox(N_new, f'contribution {k+1} text', x + 274320, y + 868680, cw - 548640, 1051560,
            [para([run(body, 14, color='tx2')])])
banner(N_new, 'not claimed banner', L, 5669280, W, 548640,
       [para([run('Not claimed: ', 16, b=True, color='accent2', font='Cambria'),
              run('clinical accuracy, glaucoma or optic-neuritis diagnosis, cross-cohort generalisation, state of the art.', 16, color='bg1', font='Cambria')], sz_end=16)],
       adj=18000, pad=274320)
set_notes(N_new, "What is new here is not a headline accuracy. First, an architecture: gated bidirectional cross-attention between spatial and temporal tokens, "
    "with squeeze-and-excitation residual blocks, multi-scale temporal convolutions, learned attention pooling and positional embeddings, implemented with "
    "deterministic pooling so GPU training is reproducible. Second, a controlled benchmark that shows when fusion helps: a synthetic generator with independent "
    "knobs for the class information in each modality, compared like for like with concatenation, element-wise sum and unimodal models. Third, a reproducible "
    "real-data protocol on two public cohorts with leave-one-cohort-out validation, repeated cross-validation and permutation tests. Fourth, candid findings, "
    "including negative ones. We do not claim clinical accuracy, diagnosis of glaucoma or optic neuritis, generalisation across cohorts, or state of the art.")

# =============================================================================
# NEW: Baselines from the literature
# =============================================================================
N_base = fac.new('Baselines from the literature')
chips = [
    ('Concatenation', 'Mean-pooled streams concatenated and projected', 'tx2'),
    ('Element-wise sum', 'Pooled stream features added element-wise', 'tx2'),
    ('Unimodal', 'MRI only, and VEP / spectrum only', 'tx2'),
    ('One-way (v1)', 'Spatial tokens query temporal tokens; mean pooling, no positional encoding or SE', 'accent3'),
    ('CAFN (ours)', 'Gated bidirectional cross-attention, attention pooling, positional embeddings, SE', 'accent1'),
]
gap = 137160
cwid = (W - 4 * gap) / 5
for k, (t, d, col) in enumerate(chips):
    x = L + k * (cwid + gap)
    ours = col == 'accent1'
    card(N_base, f'baseline {t}', x, TOP, cwid, 1371600, 6000, fill=col if ours else 'bg1')
    textbox(N_base, f'baseline {t} title', x + 137160, TOP + 118872, cwid - 274320, 320040,
            [para([run(t, 16, b=True, color='bg1' if ours else col, font='Cambria')], sz_end=16)], anchor='ctr')
    textbox(N_base, f'baseline {t} text', x + 137160, TOP + 502920, cwid - 274320, 731520,
            [para([run(d, 12, color='bg1' if ours else 'tx2')], sz_end=12)])
card(N_base, 'literature table card', L, 2834640, W, 2926080, 2500)
colw = [1463040, 1920240, 2926080, 1554480, 3047695]
rows = [
    ['Study', 'Modalities', 'Data (size)', 'Task', 'Reported result (authors\' data)'],
    ['CAMF, Zhou et al. 2024', 'sMRI + fMRI (both 3D imaging)', 'COBRE + FBIRN + MPRC (620 subjects); BSNIP (442) as independent test',
     'Schizophrenia vs control', 'Acc. 74.4%, F1 0.705, MCC 0.491 (BSNIP: 67.9%, 0.684, 0.361); cross-attention > concatenation, sum, single modality'],
    ['DRIFA-Net, Dhar et al. 2025', 'Dermoscopy, cytology, CT, MRI slices; BraTS MRI', 'HAM10000 (10,015), SIPaKMeD (4,049), brain-tumour MRI (7,023), IQ-OTHNCCD (1,098), BraTS2020 (369)',
     '3–7-class classification; segmentation', 'Acc. 95.6–99.7% (D1–D4); both attention modules > either alone'],
    ['This work', '3D T1w MRI + 1D electrophysiology (heterogeneous)', 'Controlled synthetic benchmark; two OpenNeuro cohorts (95 subjects)',
     'Control vs patient; 3-class synthetic', 'See the results slides'],
]
table(N_base, 'literature table', L + 91440, 2880360, colw, rows, [777240, 777240, 640080], hdr_h=320040, sz=11, hdr_sz=12,
      fills={3: ('accent1', 15)})
textbox(N_base, 'comparability note', L, 5852160, W, 365760,
        [para([run('Numbers are those reported by the authors on their own data; they are not comparable with ours (different tasks, modalities and data). '
                   'All our baselines share the encoders, the auxiliary loss and the training recipe.', 12, i=True, color='tx2')], sz_end=12)], anchor='ctr')
set_notes(N_base, "Published cross-attention fusion studies, for example CAMF for schizophrenia from structural and functional MRI, and DRIFA-Net on five medical imaging datasets, "
    "report a single headline number, which makes it hard to tell when fusion is responsible for a gain. CAMF reports that cross-attention beats concatenation, "
    "element-wise sum and single-modality models; DRIFA-Net reports that both of its attention modules beat either alone. We therefore implement the fusion "
    "strategies used in that literature as like-for-like baselines: concatenation, element-wise sum, the two unimodal models, and our original one-way cross-attention, "
    "all sharing the same encoders, auxiliary loss and training recipe. The table shows what those studies used and reported. Their numbers are on their data and "
    "are not comparable with ours. To our knowledge no published study fuses MRI with an electrophysiological signal for visual-pathway disorders.")

# =============================================================================
# Original slide 5 -> architecture overview with fig_arch_v2
# =============================================================================
s = S[4]
set_text(txt(s, 'Two streams'), 'Two streams, one gated fusion module')
fig_card = find(s, name='fig_pipeline card')
old_pic = find(s, name='fig_pipeline')
remove(s, old_pic)
geom(fig_card, h=3703320)
ph = 3383280
pw = ph * 1776 / 840
if pw > 7635240 - 274320:
    pw = 7635240 - 274320; ph = pw * 840 / 1776
picture(s, 'fig_arch_v2', open(ARCH, 'rb').read(), L + (7635240 - pw) / 2, TOP + (3703320 - ph) / 2, pw, ph,
        'CAFN data flow: MRI and signal encoders, gated bidirectional cross-attention, attention pool, fused and auxiliary heads')
cards = [(find(s, name='pipeline note 1'), 'Queries badge', 'Queries', 'Tokens', '64 + 16 tokens, d = 128'),
         (find(s, name='pipeline note 2'), 'Keys and values badge', 'Keys and values', 'Both directions', 'Each stream queries the other, gated'),
         (find(s, name='pipeline note 3'), 'Three heads badge', 'Three heads', 'Three heads', 'Fused plus two auxiliary heads, λ = 0.3')]
sh_list = list(s.shapes)
for k, (c, bname, oldtitle, newtitle, newtext) in enumerate(cards):
    y = TOP + k * 1280160
    body = [x for x in s.shapes if x.has_text_frame and x.top > c.top and x.top < c.top + c.height and x.left == 8659368 and x.height == 594360][0]
    t = txt(s, oldtitle)
    badge_sh = find(s, name=bname)
    icon_sh = [x for x in s.shapes if x.name == bname.replace(' badge', ' icon')][0]
    geom(c, y=y, h=1143000)
    geom(badge_sh, y=y + 137160)
    geom(icon_sh, y=y + 137160 + 125730)
    geom(t, y=y + 137160)
    set_text(t, newtitle)
    geom(body, y=y + 685800, h=411480)
    set_text(body, newtext)
banner(s, 'gate banner', L, 5212080, W, 1005840,
       [para([run('Gates start half open (σ(0) = 0.5) and the attention pool starts as mean pooling, so training begins from a neutral state and learns how much to cross-attend.', 18, color='bg1', font='Cambria')], sz_end=18)])
set_notes(s, "Here is the end-to-end data flow. The preprocessed MRI goes through a 3D residual encoder with squeeze-and-excitation blocks and is pooled to 64 spatial tokens "
    "with a learned 3D positional embedding. The 1D signal, a VEP or a power spectrum, goes through a multi-scale convolutional encoder and is pooled to 16 temporal tokens "
    "with a sinusoidal positional encoding. The fusion module is gated and bidirectional: temporal tokens read the spatial tokens, then spatial tokens read the enriched "
    "temporal tokens. A learned attention pool turns the fused spatial tokens into one vector for the fused head, and two auxiliary heads read the unimodal tokens. "
    "The gates start at sigma of zero, 0.5, and the attention pool starts as mean pooling, so training starts from a neutral state.")

# =============================================================================
# Original slide 9 -> gated bidirectional cross-attention
# =============================================================================
s = S[8]
set_text(txt(s, 'Spatial tokens query'), 'Gated bidirectional cross-attention')
remove(s, find(s, name='fig_cross_attention'))
left = find(s, name='fig_cross_attention card')
left.name = 'update equations card'
for nm in ['Per head h', 'Q = X', 'K = X', 'A = softmax', '4 heads']:
    remove(s, txt(s, nm))
remove(s, find(s, name='equations card'))
EQ = dict(font='Cambria', i=True)
# left card: the two update equations
textbox(s, 'eq header', 822960, 1508760, 6858000, 320040,
        [para([run('Each pre-norm block updates both streams, temporal first', 14, b=True, color='accent1')])])
for k, (eq, sub) in enumerate([
        ([('t′ = t + σ(g', 0), ('t', 1), (') · CA(LN(t), s)', 0)],
         [('then  t′ ← t′ + FFN', 0), ('t', 1), ('(LN(t′))      temporal tokens read the spatial tokens', 0)]),
        ([('s′ = s + σ(g', 0), ('s', 1), (') · CA(LN(s), t′)', 0)],
         [('then  s′ ← s′ + FFN', 0), ('s', 1), ('(LN(s′))      spatial tokens read the enriched temporal tokens', 0)])]):
    y = 2011680 + k * 1188720
    shape(s, f'eq step {k+1}', 822960, y, 457200, 457200, 'ellipse', fill='accent2' if k == 0 else 'accent3', shadow=False,
          paras=[para([run(str(k + 1), 20, b=True, color='bg1', font='Cambria')], algn='ctr', sz_end=20)], ins=0)
    textbox(s, f'eq {k+1}', 1508760, y - 45720, 6309360, 548640,
            [para([run(t, 26, color='tx1', sub=bool(sb), **EQ) for t, sb in eq], sz_end=26)], anchor='ctr')
    textbox(s, f'eq {k+1} sub', 1508760, y + 548640, 6309360, 320040,
            [para([run(t, 14, color='tx2', sub=bool(sb)) for t, sb in sub], sz_end=14)])
textbox(s, 'eq footnote', 822960, 4297680, 6858000, 502920,
        [para([run('CA = multi-head cross-attention; LN = layer norm; FFN = feed-forward; σ(g) is a learned scalar gate per direction.', 12, i=True, color='tx2')], sz_end=12)])
# right dark card
shape(s, 'attention card', 8229600, TOP, 3413455, 3611880, 'roundRect', 2679, 'tx2')
textbox(s, 'attn header', 8503920, 1508760, 2864815, 320040, [para([run('Attention, per head', 14, b=True, color='accent2')])])
for y, parts in [(1920240, [('Q = X', 0), ('q', 1), (' W', 0), ('Q', 1)]),
                 (2468880, [('K = X', 0), ('kv', 1), (' W', 0), ('K', 1), (',  V = X', 0), ('kv', 1), (' W', 0), ('V', 1)]),
                 (3017520, [('A = softmax(Q Kᵀ / √d', 0), ('h', 1), (')', 0)])]:
    textbox(s, 'attn eq', 8503920, y, 2864815, 457200, [para([run(t, 20, color='bg1', sub=bool(sb), **EQ) for t, sb in parts], sz_end=20)], anchor='ctr')
textbox(s, 'attn text', 8503920, 3611880, 2864815, 1143000,
        [para([run('4 heads, d_h = d / H = 32. Gates g_s and g_t are learned scalars initialised at 0, so σ(0) = 0.5.', 14, color='bg2')])])
# bottom note cards
set_text(txt(s, '64 queries'), 'Temporal queries')
set_text(txt(s, 'one per spatial token'), '16 temporal tokens query the 64 spatial tokens')
set_text(txt(s, '16 keys and values'), 'Spatial queries')
set_text(txt(s, 'temporal tokens, nominally'), '64 spatial tokens query the enriched temporal tokens')
set_text(txt(s, 'Fused tokens (B'), 'Attention pool')
set_text(txt(s, 'plus the attention map'), 'learned pool; both attention maps returned')
set_notes(s, "This is the heart of CAFN. In the original design only spatial tokens queried temporal tokens. Now both streams are updated in each pre-norm block, temporal first, "
    "so that the spatial read consumes the enriched temporal tokens. Equation one: the temporal tokens t become t plus sigma of g sub t times cross-attention with queries from "
    "t and keys and values from the spatial tokens s, followed by a feed-forward residual. Equation two: the spatial tokens s become s plus sigma of g sub s times "
    "cross-attention with queries from s and keys and values from the updated temporal tokens. The gates g are learned scalars initialised at zero, so sigma of zero is 0.5 and "
    "both directions start half open. Attention itself is standard multi-head scaled dot-product attention with four heads of dimension 32. A learned attention pool then "
    "reduces the fused spatial tokens; its near-zero initial scorer reproduces mean pooling at the start of training. Both attention maps can be returned for plotting.")

# =============================================================================
# Original slide 8 -> network size and components
# =============================================================================
s = S[7]
set_text(txt(s, 'A 1.06 M'), 'A 1.96 M-parameter two-stream network')
remove(s, find(s, name='fig_architecture'))
comps = [('3D residual CNN (MRI)', '3D ResNet + SE (MRI)', '894,504 params · 46%', 'Stem plus three SE residual blocks, 16 → 128 channels; pooled to 4×4×4 = 64 tokens.'),
         ('1D CNN (VEP)', 'Multi-scale 1D CNN', '785,936 params · 40%', 'Strided stages with parallel kernels 3 / 9 / 19; pooled to 16 tokens.'),
         ('Cross-attention fusion', 'Gated bidirectional fusion', '265,218 params · 14%', '4 heads, pre-norm, FFN in both directions; stackable blocks.'),
         ('Three linear heads', 'Embeddings, pool, heads', '9,482 params · 0.5%', 'Positional embeddings (8,192), attention pool (129), three heads 128 → 3 (1,161).')]
for old, new, params, desc in comps:
    set_text(txt(s, old), new)
    pbox = [x for x in s.shapes if x.has_text_frame and x.left == 8945575 and abs(x.top - txt(s, new).top) < 5]
    set_text(pbox[0], params)
    dbox = [x for x in s.shapes if x.has_text_frame and x.left == 5989320 and x.height == 502920 and abs(x.top - (txt(s, new).top + 457200)) < 5]
    set_text(dbox[0], desc)
# native token-flow diagram in the left card
cx, cy, cwd = 685800, 1463040, 3886200
colw_ = (cwd - 182880) / 2
x2 = cx + colw_ + 182880
textbox(s, 'col head spatial', cx, cy, colw_, 274320, [para([run('Spatial stream', 14, b=True, color='accent3', font='Cambria')], algn='ctr', sz_end=14)], anchor='ctr')
textbox(s, 'col head temporal', x2, cy, colw_, 274320, [para([run('Temporal stream', 14, b=True, color='accent2', font='Cambria')], algn='ctr', sz_end=14)], anchor='ctr')
sp_lab = ['3D T1w MRI (1, 64³)', 'Stem + 3 SE residual blocks, 16 → 128 ch.', 'Pool 4×4×4 + 3D pos-emb: 64 × 128 tokens']
tp_lab = ['VEP or spectrum (1, 500)', 'Stem + 3 strided stages, multi-scale k = 3 / 9 / 19', 'Pool + sinusoidal enc.: 16 × 128 tokens']
for r in range(3):
    y = 1828800 + r * 868680
    for (xx, lab, col, tg) in [(cx, sp_lab[r], 'accent3', 'S'), (x2, tp_lab[r], 'accent2', 'T')]:
        shape(s, f'stream {tg}{r+1}', xx, y, colw_, 685800, 'roundRect', 8000, (col, 18), line=col, shadow=False,
              paras=[para([run(lab, 12, color='tx1')], algn='ctr', sz_end=12)], ins=64008)
        if r < 2:
            arrow(s, f'stream arrow {tg}{r+1}', xx + colw_ / 2 - 70000, y + 685800 + 21000, 140000, 140000, 'downArrow')
shape(s, 'fusion box', cx, 4572000, cwd, 685800, 'roundRect', 8000, ('accent4', 18), line='accent4', shadow=False,
      paras=[para([run('Gated bidirectional cross-attention, then learned attention pool', 12, color='tx1')], algn='ctr', sz_end=12)], ins=64008)
shape(s, 'heads box', cx, 5394960, cwd, 594360, 'roundRect', 8000, ('tx2', 12), line='tx2', shadow=False,
      paras=[para([run('Fused head + 2 auxiliary heads (128 → 3)', 12, color='tx1')], algn='ctr', sz_end=12)], ins=64008)
for xx, tg in [(cx + colw_ / 2, 'L'), (x2 + colw_ / 2, 'R')]:
    arrow(s, f'to fusion {tg}', xx - 70000, 4318000, 140000, 227000, 'downArrow')
arrow(s, 'to heads', cx + cwd / 2 - 70000, 5290000, 140000, 100000, 'downArrow')
set_notes(s, "The default three-class CAFN has 1,955,140 trainable parameters. The 3D spatial encoder holds 894,504 of them, 46 percent: a stem and three squeeze-and-excitation "
    "residual blocks from 16 to 128 channels, pooled to 4 by 4 by 4, which gives the 64 spatial tokens. The multi-scale 1D encoder holds 785,936, 40 percent, with parallel "
    "convolutions of kernel 3, 9 and 19 so that both a narrow peak and a broad waveform complex are visible. The gated bidirectional fusion holds 265,218, about 14 percent. "
    "Positional embeddings, the attention pool and the three heads together are under 10,000 parameters. Spatial and temporal tokens are both 128-dimensional.")

# =============================================================================
# Original slide 10 -> loss and training
# =============================================================================
s = S[9]
remove(s, find(s, name='fig_training_flow'))
fc = find(s, name='fig_training_flow card')
r1 = ['Dataset: synthetic or real cohort', 'Seeded, stratified split', 'DataLoader, batch 8', 'CAFN forward: fused + 2 aux logits',
      'Multi-task loss L_f + 0.3 (L_s + L_t)', 'Backward + AdamW update']
flow_row(s, 'train', 685800, 1508760, 7406640, 960120, r1, ['tx2', 'tx2', 'tx2', 'accent4', 'accent4', 'tx2'], gap=182880, sz=11)
textbox(s, 'epoch note', 685800, 2514600, 7406640, 228600, [para([run('repeat for each epoch', 12, i=True, color='tx2')], algn='ctr', sz_end=12)], anchor='ctr')
r2 = ['Per-epoch validation (loss, balanced accuracy)', 'Checkpoint: best validation balanced accuracy', 'Held-out test, Grad-CAM and attention maps']
flow_row(s, 'eval', 685800, 2834640, 7406640, 960120, r2, ['tx2', 'tx2', 'accent5'], gap=228600, sz=12)
set_text(txt(s, 'Cross-entropy on'),
         'Cross-entropy on the fused, spatial and temporal heads. The auxiliary terms stop the fused head from ignoring one modality. Optional: class weights, label smoothing, modality dropout.')
set_text(txt(s, 'Adam'), 'AdamW')
set_text(txt(s, 'optimiser, lr'), 'optimiser, lr 1e-3')
set_text(txt(s, '20'), 'Cosine')
set_text(txt(s, 'epochs'), 'LR schedule, 10% warmup')
set_text(txt(s, '80 / 20'), 'Bal. acc.')
set_text(txt(s, 'split, seed 42'), 'checkpoint metric (validation)')
set_notes(s, "Training uses three cross-entropy losses: one on the fused head and one on each unimodal head. The total is the fused loss plus 0.3 times the sum of the two auxiliary losses. "
    "The auxiliary terms give each encoder its own gradient and discourage the fused head from relying on one modality only. Optional extras are class weights, label smoothing, "
    "modality dropout and a Kendall-style uncertainty loss. The defaults are AdamW at learning rate 1e-3 and weight decay 1e-4, batch size 8, linear warmup then cosine decay. "
    "Importantly, the restored checkpoint is the one with the best validation balanced accuracy, because selecting on validation loss could restore an untrained epoch-1 model on short runs. "
    "Splits and epoch counts differ per study and are given on the study slides.")

# =============================================================================
# Original slide 11 -> explanations
# =============================================================================
s = S[10]
set_text(txt(s, 'Where the 3D encoder'), 'Where the 3D encoder responds (Grad-CAM, Grad-CAM++) and which tokens each stream attends to, per head, in both directions.')
set_notes(s, "Two explainability outputs are built in. 3D Grad-CAM, and Grad-CAM++, hook the last residual block, back-propagate the fused logit, weight the channels by their gradients, "
    "apply a ReLU and upsample to the volume size. The attention maps are returned per head for both directions of the fusion; the panel B diagram shows the spatial-to-temporal map. "
    "These are inspection aids, not clinical evidence: attention weights are not automatically faithful explanations and saliency maps can be insensitive to model parameters. "
    "A planned check is to compare them against clinician-marked regions, and to anchor temporal tokens to VEP landmarks; that has not been done.")

# =============================================================================
# NEW: Two studies
# =============================================================================
N_two = fac.new('Two studies, two questions')
cw2 = 5410048
cardh = 4023360
studies = [
    (L, 'Study A · synthetic benchmark', 'accent3', 'Add baselines', 'When does fusion help?', [
        '3 classes: Healthy / Glaucoma / ON-CVI',
        'Knobs mri_strength and vep_strength in [0, 1] set the class information in each modality',
        '144 / 48 / 120 train / val / test; 32³ MRI, 500-point VEP',
        'Main setting MRI 1.0 : VEP 0.3, chosen from a pilot on the unimodal models only',
        '12 configurations, at least 5 seeds each; checkpoint by validation balanced accuracy; all numbers on the held-out test set',
        'Sweep over VEP informativeness: MRI-only, VEP-only, concatenation, CAFN']),
    (6233008, 'Study B · real cohorts', 'accent4', 'Train properly', 'Does the pipeline work on real data?', [
        'ds003568: MDD vs healthy adolescents, MEG: 18 controls, 29 patients',
        'ds007181: zoster-associated neuralgia vs healthy, PSG EEG: 27 controls, 21 patients',
        'Label: pan-disorder patient vs control; not eye disease. The 1-D stream is the log power spectrum of the resting recording',
        'Pooled model and leave-one-cohort-out; 5-fold repeated CV on all 95 subjects',
        'Bootstrap CIs, within-cohort label-permutation tests, classical baselines']),
]
for (x, title, col, icon, q, bullets) in studies:
    card(N_two, title, x, TOP, cw2, cardh, 2500)
    badge(N_two, title, x + 274320, TOP + 182880, 548640, col, ICON[icon])
    textbox(N_two, title + ' title', x + 1005840, TOP + 182880, cw2 - 1280160, 320040,
            [para([run(title, 18, b=True, color='tx1', font='Cambria')], sz_end=18)], anchor='ctr')
    textbox(N_two, title + ' question', x + 1005840, TOP + 502920, cw2 - 1280160, 228600,
            [para([run(q, 14, i=True, color=col)], sz_end=14)], anchor='ctr')
    textbox(N_two, title + ' bullets', x + 274320, TOP + 914400, cw2 - 548640, cardh - 1005840,
            [para([run(b, 14, color='tx1')], bullet=True, spc_aft=5) for b in bullets])
banner(N_two, 'neither clinical banner', L, 5486400, W, 731520,
       [para([run('Neither study is clinical. ', 16, b=True, color='accent2', font='Cambria'),
              run('A asks when fusion helps under a known ground truth; B tests the whole pipeline on real recordings that are not eye-disease data.', 16, color='bg1', font='Cambria')], sz_end=16)],
       adj=12000, pad=274320)
set_notes(N_two, "The project has two studies. Study A is a controlled synthetic benchmark. The generator has two knobs, mri_strength and vep_strength between zero and one, that scale "
    "how much class information each modality carries, so the two modalities are complementary by construction. Each run draws 144 training, 48 validation and 120 test samples, "
    "with 32-cubed MRI and 500-point VEPs. The main setting, MRI 1.0 and VEP 0.3, was fixed from a pilot that trained only the unimodal models, not tuned on CAFN. Twelve "
    "configurations are run on at least five seeds, and a sweep varies the VEP informativeness for four methods. Study B uses two public OpenNeuro cohorts: ds003568, major depressive "
    "disorder versus healthy adolescents with MEG, 18 controls and 29 patients, and ds007181, zoster-associated neuralgia versus healthy controls with polysomnographic EEG, "
    "27 controls and 21 patients. These are not eye-disease data. The label is pan-disorder patient versus control, and since no VEP is available the 1D stream is the log "
    "power spectrum of the resting recording. We evaluate a pooled model and leave-one-cohort-out, with repeated cross-validation, bootstrap intervals and permutation tests.")

# =============================================================================
# Original slide 6 -> preprocessing
# =============================================================================
s = S[5]
remove(s, find(s, name='fig_preprocessing'))
bw5 = (10545775 - 4 * 274320) / 5
mri = ['1. Load NIfTI, reorient to RAS', '2. Resample to 1 mm isotropic', '3. Otsu mask, hole filling, largest component',
       '4. Crop to foreground box, resample to 64³ (32³ for CV)', '5. Foreground z-score, add channel axis']
vep = ['1. Validate (1-D, finite, non-empty)', '2. Zero-phase band-pass, 1–100 Hz (Butterworth, 4th)', '3. Zero-phase 50 Hz notch (Q = 30)',
       '4. Linear resample to 500 samples', '5. Z-score, add channel axis']
spec = ['1. Resting MEG / EEG recording', '2. Channel-averaged log10 Welch PSD, 1–45 Hz', '3. Interpolate to 500 points', '4. Z-score; no time-domain filtering']
for yl, yb, lab, col, lst, tg in [(1463040, 1783080, 'MRI stream (src/preprocessing/mri_transforms.py)', 'accent3', mri, 'mri'),
                                  (2788920, 3108960, 'VEP stream (src/preprocessing/signal_cleaner.py)', 'accent2', vep, 'vep'),
                                  (4114800, 4434840, 'Spectrum mode, real cohorts (resting MEG / EEG instead of a VEP)', 'accent4', spec, 'spec')]:
    textbox(s, f'{tg} label', 822960, yl, 10000000, 274320, [para([run(lab, 16, b=True, color=col, font='Cambria')], sz_end=16)], anchor='ctr')
    flow_row(s, tg, 822960, yb, 10545775, 914400, lst, col, gap=274320, sz=12, bw=bw5)
set_text(txt(s, 'Output tensors'), 'Output tensors  MRI (1, 64, 64, 64)   ·   signal (1, 500)   ·   the Otsu mask is a stand-in, not validated skull-stripping')
set_notes(s, "Both modalities are brought to fixed shapes before the network. The MRI is loaded with NiBabel, reoriented to RAS, resampled to 1 millimetre isotropic and masked with an "
    "Otsu threshold with hole filling and largest-component clean-up. For real whole-head scans the volume is cropped to the foreground bounding box and resampled to 64 cubed, "
    "or 32 cubed for the cross-validation study, then z-scored over the foreground. The old centre crop saw only a sliver of a real head. The VEP is validated, band-passed 1 to 100 hertz, "
    "notch filtered at 50 hertz, both forward and backward so the effective response is zero-phase, resampled to 500 samples and z-scored. For the real cohorts, which contain resting MEG or "
    "EEG and not VEPs, each recording is reduced to the channel-averaged log10 Welch power spectrum from 1 to 45 hertz, interpolated to 500 points and z-scored, with no time-domain filtering. "
    "Neither mask is a substitute for validated skull-stripping.")

# =============================================================================
# Original slide 7 (zero-phase) : text still true; small notes
# =============================================================================
S7 = S[6]

# =============================================================================
# Original slide 4 : three classes (synthetic) - add scope in footnote
# =============================================================================
s = S[3]
set_text(txt(s, "Values are the synthetic"),
         "Synthetic generator defaults (Study A scales them with the two informativeness knobs), chosen to follow the clinical pattern of delayed P100 in optic neuritis; they are not measurements from patients.")

# =============================================================================
# Original slide 12 : synthetic data
# =============================================================================
s = S[11]
set_text(txt(s, 'Synthetic data stands'), 'Synthetic data: a controlled generator')
set_notes(s, "Because labelled paired data are unavailable, a seeded generator writes class-conditioned MRI volumes and VEP records through the same NIfTI and CSV interface as real data. "
    "On the left, the volume is Gaussian noise around 100, with an atrophy block for glaucoma and a lesion block for optic neuritis or CVI. On the right, the VEP is an N75-P100-N145 complex "
    "with drift, mains interference and noise; raw on the left, conditioned on the right. The pictures show the default generator. Two parameters, mri_strength and vep_strength, scale how far "
    "the lesion contrast and the VEP latency and amplitude move from the healthy class; the controlled benchmark uses them. With the default, unrestricted generator the VEP alone is almost perfectly "
    "separable, which is why we needed the knobs.")

# =============================================================================
# NEW: One command, one seed
# =============================================================================
N_cmd = fac.new('One command, one seed')
shape(N_cmd, 'command pill', L, TOP, 3200400, 594360, 'roundRect', 49412, 'tx1', shadow=False,
      paras=[para([run('./run_all.sh', 22, b=True, color='accent2', font='Courier New')], algn='ctr', sz_end=22)])
textbox(N_cmd, 'command text', L + 3429000, TOP, W - 3429000, 594360,
        [para([run('download → preprocess → train → validate → test. Safe to interrupt and re-run.', 18, color='tx2')], sz_end=18)], anchor='ctr')
steps = [('1', 'Environment and self-check', 'Virtualenv, GPU check, pre-flight tests of the loaders and runner', 'accent1'),
         ('2', 'Download', 'Resumable, size-verified downloads of the two cohorts (~68 GB)', 'accent3'),
         ('3', 'Preprocess', 'T1w to 64³ arrays; resting MEG / EEG to 500-point log-PSD; manifest', 'accent2'),
         ('4', 'Train, validate, test', 'Pooled, then leave-one-cohort-out; early stopping', 'accent5'),
         ('5', 'Summary', 'Tables, ROC plot and an environment snapshot', 'accent4')]
sw = 2072579
for k, (n, t, d, col) in enumerate(steps):
    x = L + k * (sw + 182880)
    y = 2103120
    card(N_cmd, f'step {n}', x, y, sw, 2286000, 4412)
    textbox(N_cmd, f'step {n} number', x + 182880, y + 137160, sw - 365760, 502920, [para([run(n, 32, b=True, color=col, font='Cambria')], sz_end=32)], anchor='ctr')
    textbox(N_cmd, f'step {n} title', x + 182880, y + 685800, sw - 365760, 548640, [para([run(t, 16, b=True, color='tx1', font='Cambria')], sz_end=16)], anchor='t')
    textbox(N_cmd, f'step {n} text', x + 182880, y + 1280160, sw - 365760, 914400, [para([run(d, 13, color='tx2')], sz_end=13)])
infos = [('SEED = 42', 'accent2', 'Repeat k uses seed 42 + k. Deterministic algorithms are requested and an environment snapshot is saved with every run.'),
         ('RTX 3050, 4 GB', 'accent3', 'NVIDIA laptop GPU, thermally throttled during long runs; the model needs under 1 GB at batch size 8. An Apple M3 Max was also used.'),
         ('Resumable', 'accent4', 'Downloads resume and are size-verified, preprocessed subjects are cached, finished training runs are skipped on re-run.')]
iw = (W - 2 * 182880) / 3
for k, (v, col, d) in enumerate(infos):
    x = L + k * (iw + 182880)
    card(N_cmd, f'info {k+1}', x, 4572000, iw, 1645920, 5000)
    textbox(N_cmd, f'info {k+1} value', x + 228600, 4663440, iw - 457200, 457200, [para([run(v, 26, b=True, color=col, font='Cambria')], sz_end=26)], anchor='ctr')
    textbox(N_cmd, f'info {k+1} text', x + 228600, 5166360, iw - 457200, 960120, [para([run(d, 14, color='tx2')])])
set_notes(N_cmd, "Everything for the real-cohort study runs from one command, run_all.sh, and is safe to interrupt and re-run. It creates the environment and runs pre-flight tests, downloads the two "
    "cohorts with resumable, size-verified downloads, preprocesses the T1-weighted scans and the resting recordings, trains and validates the pooled and leave-one-cohort-out models, and writes a "
    "summary. All randomness derives from one constant seed, 42, with repeat k using 42 plus k; deterministic algorithms are requested and the environment is recorded. Neural models were trained "
    "on an NVIDIA RTX 3050 Laptop GPU with 4 gigabytes, which throttled thermally during long runs; the model needs under one gigabyte at batch size 8. An Apple M3 Max was also used.")

# =============================================================================
# Original slide 15 -> compute
# =============================================================================
s = S[14]
set_text(txt(s, 'Compute sits'), 'Compute sits in the 3D spatial encoder')
remove(s, find(s, name='fig_params'))
cap = txt(s, 'Left: trainable')
geom(cap, y=4160520, h=320040)
set_text(cap, 'Left: share of the 1,955,140 trainable parameters. Right: share of the 2.60 G multiply–accumulates per sample at 64³.')
comp_rows = [('Spatial encoder', 'accent3', 45.75, '894,504 · 46%', 97.03, '2.52 G · 97%'),
             ('Temporal encoder', 'accent2', 40.20, '785,936 · 40%', 2.53, '0.07 G · 2.5%'),
             ('Gated fusion', 'accent4', 13.57, '265,218 · 14%', 0.42, '0.011 G · 0.4%'),
             ('Embeddings, pool, heads', 'tx2', 0.48, '9,482 · 0.5%', 0.0, '≈ 0')]
for px, pw_, hd, idx_p in [(731520, 3566160, 'Trainable parameters', 2), (4480560, 3566160, 'MACs per sample (64³)', 4)]:
    textbox(s, f'panel head {idx_p}', px, 1463040, pw_, 274320, [para([run(hd, 14, b=True, color='tx1', font='Cambria')], sz_end=14)], anchor='ctr')
    for k, row in enumerate(comp_rows):
        y = 1874520 + k * 594360
        name_, col, sp_, lp_, sm_, lm_ = row
        share, lab = (sp_, lp_) if idx_p == 2 else (sm_, lm_)
        textbox(s, f'bar label {idx_p}.{k}', px, y, pw_, 228600,
                [para([run(name_ + '  ', 12, color='tx2'), run(lab, 12, b=True, color='tx1')], sz_end=12)], anchor='ctr')
        shape(s, f'bar {idx_p}.{k}', px, y + 250000, max(30000, 3200000 * share / 100), 170000, 'rect', fill=col, shadow=False)
set_text(txt(s, '1.06 M'), '1,955,140')
set_text(txt(s, 'trainable parameters'), 'trainable parameters')
set_text(txt(s, '2.53 G'), '2.60 G')
set_text(txt(s, 'multiply–accumulates'), 'MACs per sample at 64³')
set_text(txt(s, '82%'), '97%')
set_text(txt(s, 'of MACs in residual'), 'of MACs in the spatial encoder')
set_text(txt(s, 'Residual block 1 alone'),
         'The 3D spatial encoder holds 46% of the parameters but 97% of the MACs (2.60\u00a0G at 64³, 0.39\u00a0G at 32³). The gated bidirectional fusion holds 14% of the parameters and under 0.5% of the compute, so speed-ups should target the 3D encoder.')
set_notes(s, "On cost: 1,955,140 parameters and 2.60 billion multiply-accumulates per sample at 64 cubed, counted analytically from the layer definitions; at 32 cubed it is 0.39 billion. "
    "The 3D spatial encoder holds 46 percent of the parameters and 97 percent of the compute. The multi-scale temporal encoder holds 40 percent of the parameters but only about 2.5 percent of the compute, "
    "and the gated bidirectional fusion holds 14 percent of the parameters and under half a percent of the compute. So if speed matters, the 3D encoder is the place to optimise, not the fusion module. "
    "The whole model needs under one gigabyte of GPU memory at batch size 8.")

# =============================================================================
# Original slide 16 -> tests
# =============================================================================
s = S[15]
set_text(txt(s, '86 tests'), '340 tests pass on two platforms')
set_text(find(s, prefix='86', nth=0), '340')
set_text(txt(s, 'tests: 78 unit'), 'tests, unit and integration; all pass on macOS, Python 3.14')
set_text(txt(s, '95.1%'), 'Ubuntu')
set_text(txt(s, 'statement coverage (521'), '22.04 (WSL2), Python 3.12, CUDA: one test skipped')
set_text(txt(s, '8 of 12'), '3 bugs')
set_text(txt(s, 'modules fully covered'), 'latent defects exposed by the second platform, all fixed')
remove(s, find(s, name='fig_tests_coverage'))
remove(s, txt(s, 'Left: test functions'))
rc = find(s, name='fig_tests_coverage card')
textbox(s, 'bugs header', 4251960, 1508760, 7160000, 365760, [para([run('What the second platform exposed', 20, b=True, color='tx1', font='Cambria')], sz_end=20)], anchor='ctr')
bugs = ['A type annotation in utils.py raised a NameError on Python 3.13 and earlier',
        'The --eval-only path had a device mismatch on CUDA',
        'A unit test whose assertion was mathematically vacuous']
for k, b in enumerate(bugs):
    y = 2148840 + k * 777240
    shape(s, f'bug {k+1} badge', 4251960, y, 457200, 457200, 'ellipse', fill='accent6', shadow=False,
          paras=[para([run(str(k + 1), 18, b=True, color='bg1', font='Cambria')], algn='ctr', sz_end=18)], ins=0)
    textbox(s, f'bug {k+1}', 4937760, y - 22860, 6477000, 502920, [para([run(b, 16, color='tx1')], sz_end=16)], anchor='ctr')
textbox(s, 'also verified', 4251960, 4526280, 7160000, 960120,
        [para([run('Also verified: the CTF MEG and EDF readers on real recordings, and that the deterministic adaptive pooling matches PyTorch in forward and backward passes.', 14, i=True, color='tx2')], sz_end=14)])
set_text(txt(s, 'Counts and coverage'),
         "The suite covers unit and integration tests, including end-to-end runs of the command-line tools, the cohort loaders and the experiment runner. Tests check behaviour, not clinical validity.")
set_notes(s, "On verification: the repository's automated suite now contains 340 tests, unit and integration, including end-to-end runs of the command-line tools, the cohort loaders and the experiment runner. "
    "All pass on macOS with Python 3.14, and on Ubuntu 22.04 under WSL2 with Python 3.12 and CUDA, where one test is skipped. Running on a second platform exposed three latent defects the first platform "
    "had masked: a type annotation that fails on Python 3.13 and earlier, a device mismatch in the evaluation-only path on CUDA, and a unit test whose assertion was mathematically vacuous. All three were fixed. "
    "The two real-file readers, CTF MEG and EDF polysomnography, were checked on real recordings, and the deterministic adaptive pooling was verified against PyTorch's pooling in forward and backward passes. "
    "Tests check behaviour, not clinical validity.")

# =============================================================================
# Original slides 13, 14 (synthetic conditioning results) : label + notes
# =============================================================================
s = S[12]
set_text(txt(s, 'trials within 4 ms'), 'synthetic trials within 4 ms (one sample) of the ground truth')
set_notes(s, S[12].notes_slide.notes_text_frame.text + " These are synthetic records; no patient data are involved.")

# =============================================================================
# Result slides (phase 2: real figures and final numbers)
# =============================================================================
import io
from PIL import Image
RA = os.path.join(DOCS, 'results_assets')


def png_blob(name, crop=None):
    im = Image.open(os.path.join(RA, name))
    if crop:
        im = im.crop(crop)
    buf = io.BytesIO(); im.save(buf, 'PNG')
    return buf.getvalue(), im.size


def fig_slide(title, tag, tag_color, png, crop, caption, stats, banner_paras, notes, fig_w_max=None):
    """Figure card (left, 7772400 x 3474720), three stat cards (right), banner (bottom)."""
    r = fac.new(title)
    shape(r, 'result tag', 8595360, 502920, 3047695, 411480, 'roundRect', 49412, tag_color, shadow=False,
          paras=[para([run(tag, 14, b=True, color='bg1')], algn='ctr', sz_end=14)], ins=45720)
    CH = 3474720
    card(r, 'fig card', L, TOP, 7772400, CH, 2857)
    blob, (pw_, ph_) = png_blob(png, crop)
    aw, ah = (fig_w_max or 7589520), CH - 91440 - 365760
    sc = min(aw / pw_, ah / ph_)
    w, h = pw_ * sc, ph_ * sc
    fx = L + 91440 + (0 if fig_w_max else (7589520 - w) / 2)
    pic = picture(r, 'result figure', blob, fx, TOP + 91440, w, h, caption)
    textbox(r, 'fig caption', L + 182880, TOP + CH - 365760, 7772400 - 365760, 320040,
            [para([run(caption, 12, i=True, color='tx2')], sz_end=12)], anchor='ctr')
    for k, (val, col, lab) in enumerate(stats):
        y = TOP + k * (1097280 + 91440)
        card(r, f'stat{k+1} card', 8595360, y, 3047695, 1097280, 8000)
        textbox(r, f'stat{k+1} value', 8778240, y + 82296, 2682240, 457200, [para([run(val, 28, b=True, color=col, font='Cambria')], sz_end=28)], anchor='ctr')
        textbox(r, f'stat{k+1} label', 8778240, y + 566928, 2682240, 480000, [para([run(lab, 13, color='tx2')], sz_end=13)])
    banner(r, 'takeaway banner', L, 5029200, W, 1188720, banner_paras, adj=7000)
    set_notes(r, notes)
    return r, (fx, w, h)


def bpara(lead, body, sz=17):
    return [para([run(lead + ' ', sz, b=True, color='accent2', font='Cambria'), run(body, sz, color='bg1', font='Cambria')], sz_end=sz)]


P_abl, _ = fig_slide('Ablation: which components matter?', 'SYNTHETIC BENCHMARK', 'accent3', 'fig_ablation.png', None,
    'Test accuracy by configuration, mean ± sd over seeds; dotted line = chance (33.3%); grey = single-modality references.',
    [('+10.7 pts', 'accent4', 'fusion over the better single modality (MRI only), p < 0.001, 5 paired seeds'),
     ('78.1%', 'accent3', 'CAFN; concatenation 77.5%, element-wise sum 79.6%: not different'),
     ('0 of 7', 'accent2', 'component ablations reliably different (range −1.7 to 0.0 points)')],
    bpara('Fusion matters; the specific mechanism does not, at this scale.',
          'CAFN is statistically indistinguishable from concatenation and element-wise sum, and removing any single component changes accuracy by at most 1.7 points (MRI 1.0 : VEP 0.3, 4–7 seeds, held-out test set of 120).'),
    "Controlled synthetic benchmark, main setting MRI strength 1.0 and VEP strength 0.3, three classes, held-out test set of 120, chance 33.3 percent, 4 to 7 seeds per configuration. "
    "The two single-modality references reach 68.0 percent (MRI only) and 64.0 percent (VEP only), whereas every fusion model is at 77.5 to 79.6 percent. The full CAFN, 78.1 percent, beats the better single modality by 10.7 percentage points, "
    "paired over 5 seeds, p below 0.001. But CAFN is indistinguishable from concatenation, 77.5 percent, p of 0.63, and from element-wise sum, 79.6 percent, p of 0.47, and none of the seven component ablations is reliably different, "
    "ranging from minus 1.7 to 0.0 points. So the honest reading is that fusion matters, and the specific cross-attention mechanism and its components do not make a measurable difference at this scale. This is synthetic data, not clinical performance.")

P_sweep, (sfx, sw_, sh_) = fig_slide('Informativeness sweep', 'SYNTHETIC BENCHMARK', 'accent3', 'fig_sweep.png', None,
    'Test accuracy vs VEP informativeness (MRI fixed at 1.0); 2 seeds per point, bars = sd. Cat. = concatenation.',
    [('+13 pts', 'accent4', 'CAFN over the best single modality at VEP 0.3, the largest gain'),
     ('+1 pt', 'accent3', 'at VEP 1.0: fusion converges to the dominant modality'),
     ('CAFN ≈ concat.', 'accent2', 'the two track each other at every point; 2 seeds: trends only')],
    bpara('Fusion never hurt in this sweep.',
          'It helps most when both streams carry independent evidence and converges to the dominant stream otherwise; CAFN tracks concatenation throughout. With only 2 seeds per point the curves are a trend, not precise estimates.'),
    "Sweep of the VEP informativeness knob with the MRI fixed at 1.0, for MRI-only, VEP-only, concatenation and CAFN, with only 2 seeds per point. The VEP-only model rises from 45 percent to 97 percent while the MRI-only model stays at 63 to 70 percent. "
    "The gain of the fused models over the better single modality is largest at intermediate informativeness, about 13 points for CAFN and 12 for concatenation at VEP 0.3, and vanishes when one modality dominates, about 1 to 2 points at VEP 1.0. "
    "Fusion never hurt, and CAFN and concatenation track each other throughout. Because there are only two seeds per point the intervals are wide, so read this as a trend. Synthetic data only.")
x0 = sfx + sw_ + 182880
textbox(P_sweep, 'gain table header', x0, TOP + 91440, L + 7772400 - x0 - 137160, 420000,
        [para([run('Gain over best single modality (points)', 14, b=True, color='tx1', font='Cambria')], sz_end=14)], anchor='ctr')
table(P_sweep, 'gain table', x0, TOP + 548640, [(L + 7772400 - x0 - 137160) * 0.34, (L + 7772400 - x0 - 137160) * 0.33, (L + 7772400 - x0 - 137160) * 0.33],
      [['VEP', 'CAFN', 'Cat.'], ['0.15', '+5', '+7'], ['0.3', '+13', '+12'], ['0.5', '+5', '+5'], ['1.0', '+1', '+2']],
      411480, hdr_h=365760, sz=14, hdr_sz=12, first_bold=False, algns=['ctr', 'ctr', 'ctr'])

P_neural, _ = fig_slide('Neural vs classical, real cohorts', 'REAL COHORTS · 95 subjects', 'accent4', 'fig_real_cv.png', (0, 0, 1280, 600),
    'Out-of-fold ROC-AUC (mean over repeats), 95% bootstrap CI; dotted = chance. Grey classical, blue neural.',
    [('0.62', 'accent4', 'spectrum-only network AUC, CI [0.50, 0.72]: the best model'),
     ('≈ 0.5', 'accent6', 'CAFN 0.51, concatenation 0.53, MRI-only 0.48: at chance'),
     ('0.567', 'accent3', 'band-power LR AUC; p = 0.005, Bonferroni 0.025; CI includes 0.5')],
    bpara('Weak at best.',
          'Adding the MRI stream did not help on these cohorts and the fused networks are at chance; only a weak spectral signal remains, and its evidence is at most suggestive. Neural models were cross-validated once, with no permutation test.'),
    "Out-of-fold performance on all 95 subjects. The best neural model is the spectrum-only network, AUC 0.62 with a bootstrap interval of 0.50 to 0.72. The full CAFN, 0.51, concatenation, 0.53, and the MRI-only network, 0.48, are at chance, "
    "so adding the MRI stream did not help. Among the classical models, band-power logistic regression has AUC 0.567 with a within-cohort permutation p of 0.005; with five classical models tested only this one survives a Bonferroni correction, adjusted p of 0.025, "
    "and its bootstrap interval, 0.45 to 0.68, still includes 0.5, so the evidence is at best suggestive. MRI-derived features are not significant. Neural models were cross-validated once and have no permutation test. These are real recordings, but not eye-disease data.")

N_curves, (cfx, cw_, ch_) = fig_slide('The model memorises', 'REAL COHORTS · pooled CAFN', 'accent4', 'fig_real_curves.png', None,
    'Pooled CAFN on the real cohorts: training loss and accuracy vs validation (loss left, accuracy right; dotted line = chance).',
    [('→ 0', 'accent6', 'training loss within about 30 epochs'),
     ('≈ 0.5', 'accent2', 'validation accuracy throughout: chance level'),
     ('~67', 'accent3', 'training subjects for 1.96 M parameters')],
    bpara('Memorisation, not learning.',
          'With roughly 67 training subjects, two unrelated disorders under one label and different sensors, the model fits its training set and learns nothing that transfers. The bottleneck is data, not architecture.'),
    "The pooled CAFN on the real cohorts. The training loss falls to almost zero within about 30 epochs and training accuracy reaches 1.0, while validation accuracy stays near 0.5 and validation loss rises: the 1.96 million parameter model memorises the roughly 67 training subjects. "
    "Together with the cohort-identity shortcut, different sensors and a pan-disorder label, this is why the neural fusion results are at chance. The conclusion is that the bottleneck is data, not architecture.")

P_concl = fac.new('Final conclusions')
card(P_concl, 'conclusion card', L, TOP, W, 3840480, 2500)
concl = [('Fusion helps when the evidence is complementary', 'accent3',
          'In the synthetic benchmark every fusion model beats the better single modality by about 10.7 points (paired over 5 seeds, p < 0.001); the gain is largest at intermediate informativeness and vanishes when one modality dominates.'),
         ('Cross-attention is not shown to beat simpler fusion', 'accent2',
          'CAFN 78.1% vs concatenation 77.5% and element-wise sum 79.6%: indistinguishable. None of the 7 component ablations is reliably different (−1.7 to 0.0 points).'),
         ('Real cohorts: chance-level fusion, weak spectral signal', 'accent4',
          'Neural fusion is at chance (CAFN AUC 0.51); spectrum-only network 0.62 [0.50, 0.72]; band-power LR 0.57 (Bonferroni p = 0.025, CI includes 0.5); nothing transfers across cohorts. The next step is data, not architecture.')]
for k, (t, col, body) in enumerate(concl):
    y = TOP + 137160 + k * 1234440
    shape(P_concl, f'conclusion {k+1} badge', L + 228600, y + 22860, 502920, 502920, 'ellipse', fill=col, shadow=False,
          paras=[para([run(str(k + 1), 20, b=True, color='bg1', font='Cambria')], algn='ctr', sz_end=20)], ins=0)
    textbox(P_concl, f'conclusion {k+1} title', L + 960120, y, W - 1188720, 411480, [para([run(t, 20, b=True, color='tx1', font='Cambria')], sz_end=20)], anchor='ctr')
    textbox(P_concl, f'conclusion {k+1} text', L + 960120, y + 438912, W - 1188720, 740000, [para([run(body, 15, color='tx2')], sz_end=15)])
banner(P_concl, 'takeaway banner', L, 5349240, W, 868680,
       [para([run('Not claimed: ', 16, b=True, color='accent2', font='Cambria'),
              run('clinical accuracy, diagnosis of glaucoma, optic neuritis or CVI, cross-cohort generalisation, or that cross-attention beats simpler fusion.', 16, color='bg1', font='Cambria')], sz_end=16)], adj=10000)
set_notes(P_concl, "Three conclusions. One: fusion helps when both modalities carry independent evidence. In the controlled synthetic benchmark every fusion model beats the better single modality by about 10.7 points, paired over 5 seeds, p below 0.001, and the sweep shows the gain is largest at intermediate informativeness and vanishes when one modality dominates. "
    "Two: cross-attention is not shown to beat simpler fusion. CAFN, 78.1 percent, is indistinguishable from concatenation, 77.5, and element-wise sum, 79.6, and none of the seven component ablations is reliably different. "
    "Three: on the real cohorts neural fusion is at chance, CAFN AUC 0.51, only a weak spectral signal exists, with the spectrum-only network at 0.62 and band-power logistic regression at 0.57, which survives Bonferroni correction but whose interval includes 0.5, and nothing transfers across cohorts. "
    "The next step is data, specifically a labelled cohort with paired MRI and VEP, not a new architecture. We do not claim clinical accuracy, diagnosis of glaucoma, optic neuritis or CVI, cross-cohort generalisation, or that cross-attention beats simpler fusion.")

# =============================================================================
# NEW: Real-cohort results so far
# =============================================================================
N_real = fac.new('Real-cohort results so far')
shape(N_real, 'real tag', 8595360, 502920, 3047695, 411480, 'roundRect', 49412, 'accent4', shadow=False,
      paras=[para([run('REAL COHORTS · 95 subjects', 14, b=True, color='bg1')], algn='ctr', sz_end=14)], ins=45720)
card(N_real, 'classical results card', L, TOP, 7772400, 4206240, 2000)
textbox(N_real, 'table header', L + 182880, TOP + 91440, 7400000, 320040,
        [para([run('Classical baselines, 5 × 5-fold CV on all 95 subjects', 16, b=True, color='tx1', font='Cambria')], sz_end=16)], anchor='ctr')
rows = [['Model', 'ROC-AUC', '95% CI', 'Perm. p'],
        ['Spectrum, logistic regression', '0.563', '0.45–0.68', '0.025'],
        ['Band power, logistic regression', '0.567', '0.45–0.68', '0.005'],
        ['Band power, random forest', '0.518', '0.36–0.60', '0.73'],
        ['MRI (8³ → PCA 20), logistic regression', '0.587', '0.43–0.67', '0.24'],
        ['MRI + band power, logistic regression', '0.588', '0.43–0.67', '0.23']]
table(N_real, 'classical table', L + 182880, TOP + 502920, [3566160, 1143000, 1417320, 1097280], rows, 457200, hdr_h=411480, sz=14, hdr_sz=14,
      first_bold=False, algns=['l', 'ctr', 'ctr', 'ctr'])
textbox(N_real, 'table caption', L + 182880, TOP + 3383280, 7400000, 731520,
        [para([run('AUC is the mean over repeats; the CI is a bootstrap over subjects and the p-value a within-cohort label-permutation test, both for repeat 0. '
                   'Only the two spectrum-based logistic models reach p < 0.05 (AUC ≈ 0.56); after Bonferroni correction only band power survives (p = 0.025) and its CI includes 0.5.', 12, i=True, color='tx2')], sz_end=12)])
stats = [('0.45–0.52', 'accent2', 'CAFN test accuracy, pooled and LODO (balanced accuracy 0.46–0.54): at chance'),
         ('100%', 'accent6', 'cohort identity decoded from the spectrum alone: a shortcut the permutation test controls for'),
         ('0.40–0.63', 'accent3', 'classical leave-one-cohort-out AUC: no consistent transfer between cohorts')]
for k, (v, col, lab) in enumerate(stats):
    y = TOP + k * 1371600 * 1.0
    y = TOP + k * 1417320
    card(N_real, f'stat{k+1} card', 8595360, y, 3047695, 1371600, 6000)
    textbox(N_real, f'stat{k+1} value', 8778240, y + 91440, 2682240, 457200, [para([run(v, 30, b=True, color=col, font='Cambria')], sz_end=30)], anchor='ctr')
    textbox(N_real, f'stat{k+1} label', 8778240, y + 594360, 2682240, 731520, [para([run(lab, 13, color='tx2')], sz_end=13)])
banner(N_real, 'honest banner', L, 5669280, W, 548640,
       [para([run('Honest reading: ', 16, b=True, color='accent2', font='Cambria'),
              run('weak spectral signal at best, no transfer across cohorts, no evidence that fusion helps on these cohorts.', 16, color='bg1', font='Cambria')], sz_end=16)], adj=18000, pad=274320)
set_notes(N_real, "These are the real-cohort results that are already final, on 95 subjects from the two public cohorts, and they are modest. Classical models with 5 by 5-fold cross-validation reach "
    "an AUC of 0.56 to 0.59: the spectrum and band-power logistic regressions are nominally significant under a within-cohort label-permutation test, p of 0.025 and 0.005, but the "
    "AUC is only about 0.56 and only band power survives a Bonferroni correction, adjusted p of 0.025, with an interval that includes 0.5; the random forest on band powers and the MRI-based models are not significant. Leave-one-cohort-out AUC for the classical models ranges from 0.40 to 0.63: "
    "no consistent transfer. CAFN, pooled and leave-one-cohort-out over three repeats, has test accuracies of 0.45 to 0.52 and balanced accuracies of 0.46 to 0.54, which is chance. "
    "A cohort-identity probe shows the spectrum alone identifies the cohort with 100 percent accuracy, which is why the permutation test shuffles labels within each cohort. "
    "We do not claim that fusion helps on these cohorts, and the labels are pan-disorder, not eye disease.")

# =============================================================================
# Original slide 17 -> lessons
# =============================================================================
s = S[16]
set_text(txt(s, 'What the synthetic'), 'What the synthetic studies taught us')
set_text(txt(s, 'It zeroes exactly'), 'The original percentile mask zeroes exactly 60% of voxels; validated skull-stripping is still needed.')
set_text(txt(s, 'Optimise early 3D'), 'Unrestricted data is too easy')
set_text(txt(s, 'Compute concentrates'), 'VEP-only 99.0% vs MRI-only 70.8% on the default generator: it shows the pipeline trains, not that fusion helps.')
set_text(txt(s, 'of the Glaucoma atrophy'), 'of the Glaucoma atrophy block survives the 60th-percentile mask; 99.7% of the lesion block does (synthetic volumes)')
set_notes(s, "Three observations from the synthetic experiments. First, the original percentile mask is not neutral: by construction it zeroes 60 percent of voxels, and in the glaucoma volume it removes the "
    "entire atrophy block, 0 percent retained, while 99.7 percent of the optic neuritis lesion survives. Validated skull-stripping is needed before drawing conclusions about the MRI. Second, tokens are not time bins: "
    "P100 at 100 and 108 milliseconds both fall in token 0, and 140 milliseconds falls in token 1, so attention should not be read at millisecond resolution. Third, and this motivated the controlled benchmark: "
    "with the default, unrestricted generator the VEP alone is almost perfectly separable, 99.0 percent for the VEP-only head against 70.8 percent for the MRI-only head, so earlier runs that reached 97.9 percent "
    "validation accuracy and 99.0 percent held-out accuracy on 96 samples demonstrate that the pipeline trains, not that fusion helps. All of this is synthetic data.")

# =============================================================================
# NEW: Pitfalls
# =============================================================================
N_pit = fac.new('Pitfalls we found along the way')
pits = [('Cohort identity is a shortcut', 'accent6', 'warning', 'The spectrum alone identifies the cohort with 100% accuracy. Permutation tests therefore shuffle labels within each cohort, so a cohort shortcut cannot create significance.'),
        ('Checkpoint selection', 'accent2', 'Validate masking', 'Selecting on validation loss could restore an untrained epoch-1 model on short runs (smoke runs ended at 33% accuracy). Checkpoints are now chosen by validation balanced accuracy.'),
        ('A centre crop sees a sliver', 'accent3', 'Unimodal isolation', 'A centre crop to 64³ is fine for small synthetic volumes but sees only a sliver of a 170 mm head. Real scans are cropped to the foreground box, then resampled.'),
        ('Bugs hide on one platform', 'accent5', 'Verified', 'A second platform exposed an annotation error on Python 3.13 and earlier, an evaluation-only device mismatch on CUDA and a vacuous test; all fixed.')]
for k, (t, col, icon, body) in enumerate(pits):
    x = L if k % 2 == 0 else 6233008
    y = TOP if k < 2 else 3703320
    card(N_pit, f'pitfall {k+1}', x, y, cw, 2286000 if k < 2 else 2286000, 4000)
    badge(N_pit, t, x + 274320, y + 228600, 548640, col, ICON[icon])
    textbox(N_pit, f'pitfall {k+1} title', x + 1005840, y + 228600, cw - 1280160, 548640, [para([run(t, 20, b=True, color='tx1', font='Cambria')], sz_end=20)], anchor='ctr')
    textbox(N_pit, f'pitfall {k+1} text', x + 274320, y + 960120, cw - 548640, 1234440, [para([run(body, 16, color='tx2')], sz_end=16)])
set_notes(N_pit, "Four pitfalls we found and fixed. One: on the real cohorts a spectrum alone identifies the cohort with 100 percent accuracy, so a model can look good by learning cohort identity; "
    "our permutation tests shuffle labels within each cohort for that reason. Two: choosing the checkpoint by validation loss could restore an untrained epoch-1 model on short runs, and smoke and demo "
    "runs ended at 33 percent accuracy for that reason; we now select by validation balanced accuracy. Three: a centre crop to 64 cubed, enough for small synthetic volumes, sees only a sliver of a real head, "
    "so real scans are cropped to the foreground bounding box and resampled. Four: running on a second platform exposed three latent bugs.")

# =============================================================================
# Original slide 18 -> evidence
# =============================================================================
s = S[17]
shown = ['Synthetic: fusion beats the better single modality by 10.7 points (p < 0.001, 5 paired seeds)',
         'Conditioning keeps P100 within 4 ms in 149 of 150 synthetic trials; 22 dB mains reduction',
         '1,955,140 parameters, 2.60 G MACs; 340 tests pass on two platforms',
         'One-command, seeded, deterministic pipeline on two public real cohorts',
         'Real cohorts: only a weak spectral signal (band-power AUC 0.57, Bonferroni p = 0.025); fused networks at chance']
notshown = ['Diagnosis of glaucoma, optic neuritis or CVI: no eye-disease data were available',
            'That cross-attention beats concatenation or sum (78.1% vs 77.5% and 79.6%), or that any single component matters (0 of 7 ablations reliable)',
            'A benefit of fusion on real data, or generalisation across cohorts (classical LODO AUC 0.40–0.63)',
            'Clinical accuracy or state-of-the-art performance',
            'Validated skull-stripping, or faithful explanations']
set_paras(find(s, name='Text 4') if False else txt(s, 'Conditioning keeps'), [para([run(t, 15, color='tx1')], bullet=True, spc_aft=8, sz_end=15) for t in shown])
set_paras(txt(s, 'No training run'), [para([run(t, 15, color='tx1')], bullet=True, spc_aft=8, sz_end=15) for t in notshown])
set_notes(s, "To summarise the evidence honestly. What is shown: in the controlled synthetic benchmark fusion beats the better single modality by 10.7 points, p below 0.001; the conditioning chain preserves P100 latency on synthetic VEPs; "
    "model size and cost are computed; 340 tests pass on two platforms; the real-cohort pipeline is one command, seeded and deterministic; and on the two real cohorts only a weak spectral signal exists, band-power AUC 0.57, "
    "which survives Bonferroni correction with an interval that includes 0.5, while the fused networks are at chance. What is not shown: diagnosis of glaucoma, optic neuritis or CVI, because no eye-disease data were available; "
    "that cross-attention beats concatenation or sum, or that any single component matters; a benefit of fusion on real data or generalisation across cohorts; clinical accuracy or state of the art; validated skull-stripping; and faithful explanations.")

# =============================================================================
# Original slide 19 -> done / next
# =============================================================================
s = S[18]
set_text(txt(s, 'Next: a labelled'), 'What is done, what is next')
steps_new = [('1', 'Done', 'accent4', 'Train properly', 'Real-cohort protocol', 'Pooled and leave-one-cohort-out models, repeated CV and permutation tests on two public cohorts.'),
             ('2', 'Done', 'accent4', 'Add baselines', 'Baselines and ablation', 'Concatenation, sum and unimodal baselines, ablation and sweep on the synthetic benchmark.'),
             ('3', 'Next', 'accent2', 'Validate masking', 'Skull-stripping', 'Replace the Otsu mask with validated skull-stripping, for example HD-BET or SynthStrip.'),
             ('4', 'Next', 'accent2', 'Check explanations', 'Check explanations', 'Compare Grad-CAM and attention with clinician-marked regions.'),
             ('5', 'Next', 'accent2', 'Ground the tokens', 'Paired MRI + VEP cohort', 'Find or collect a labelled cohort with paired MRI and VEP for visual-pathway disorders.')]
old_titles = ['Train properly', 'Add baselines', 'Validate masking', 'Check explanations', 'Ground the tokens']
old_text = ['Patient-level splits', 'Unimodal and concatenation', 'Replace the percentile', 'Compare Grad-CAM', 'Align temporal tokens']
for k, (n, status, col, ot, nt, body) in enumerate(steps_new):
    numbox = txt(s, n, nth=0) if False else [x for x in s.shapes if x.has_text_frame and x.text_frame.text == n][0]
    set_text(numbox, status)
    set_color(numbox, col)
    set_text(txt(s, old_titles[k]), nt)
    set_text(txt(s, old_text[k]), body)
set_notes(s, "Some of what used to be future work is now done. Done: the real-cohort protocol with leave-one-cohort-out validation, repeated cross-validation and permutation tests, and the baselines, ablation "
    "and informativeness sweep on the controlled synthetic benchmark. Next: replace the Otsu foreground mask with validated skull-stripping; compare Grad-CAM and attention against clinician-marked regions; and, "
    "most importantly, obtain a labelled cohort with paired MRI and VEP for visual-pathway disorders, which is the only way to test the clinical question.")

# =============================================================================
# Original slide 20 -> closing
# =============================================================================
s = S[19]
set_text(txt(s, '1.06 M'), '1.96 M')
set_text(txt(s, 'parameters, 2.53 G'), 'parameters, 2.60 G MACs at 64³')
set_text(txt(s, '95.1%'), '≈ 50%')
set_text(txt(s, 'statement coverage'), 'CAFN test accuracy on the real cohorts: chance level')
set_text(txt(s, 'P100 trials within'), 'synthetic P100 trials within one sample')
set_text(txt(s, '22 dB'), '+10.7 pts')
set_text(txt(s, 'mains interference removed'), 'synthetic: fusion over the better single modality')
set_text(txt(s, 'CAFN is a cross-attention'),
         'CAFN is a gated bidirectional cross-attention network for 3D MRI and a 1D signal. Fusion beats the better single modality in a controlled synthetic benchmark, but cross-attention is not shown to beat simpler fusion, and on the real cohorts it is at chance. No diagnostic accuracy is claimed.')
set_notes(s, "To conclude: CAFN brings a 3D MRI and a 1D signal into a common token space and fuses them with gated bidirectional cross-attention, in a deterministic, tested and reproducible implementation. "
    "In the controlled synthetic benchmark fusion beats the better single modality by 10.7 points, but cross-attention is statistically indistinguishable from concatenation and element-wise sum and no component ablation is reliable. "
    "P100 was recovered within one sample in 149 of 150 synthetic trials. On the two real cohorts CAFN is at chance, about 50 percent accuracy, and only a weak spectral signal exists. The next step is data, not architecture. "
    "No diagnostic accuracy is claimed. Thank you; I am happy to take questions.")

# =============================================================================
# Original slide 21 -> references (two columns)
# =============================================================================
s = S[20]
set_text(txt(s, 'Selected references') if False else [x for x in s.shapes if x.is_placeholder][0], 'Selected references')
old_box = txt(s, '[1]')
old_refs = [p.text for p in old_box.text_frame.paragraphs]
new_refs = [
    '[12]  Zhou Z. et al., “An interpretable cross-attentive multi-modal MRI fusion framework for schizophrenia diagnosis,” arXiv:2404.00144, 2024.',
    '[13]  Dhar J. et al., “Multimodal fusion learning with dual attention for medical imaging,” WACV, 2025, arXiv:2412.01248.',
    '[14]  Hu J., Shen L., Sun G., “Squeeze-and-excitation networks,” CVPR, 2018.',
    '[15]  Liuzzi L. et al., “Mood induction in MDD and healthy adolescents,” OpenNeuro ds003568, v1.0.1.',
    '[16]  Li L. et al., “Structural MRI, resting-state fMRI, and PSG/EEG dataset of zoster-associated neuralgia,” OpenNeuro ds007181, v1.0.1.']
refs_all = old_refs + new_refs
remove(s, old_box)
half = 8
def ref_paras(lst):
    return [para([run(t, 13, color='tx1')], spc_aft=8, sz_end=13) for t in lst]
textbox(s, 'refs col 1', 1554480, 1508760, 4754880, 4572000, ref_paras(refs_all[:half]))
textbox(s, 'refs col 2', 6537960, 1508760, 4754880, 4572000, ref_paras(refs_all[half:]))
set_notes(s, "Backup slide. The full reference list is in the paper. New in this version, all checked against their sources: CAMF (Zhou et al., arXiv 2404.00144), DRIFA-Net (Dhar et al., WACV 2025, arXiv 2412.01248), "
    "squeeze-and-excitation networks (Hu et al., CVPR 2018) and the two OpenNeuro datasets ds003568 and ds007181. Other entries were written from known publication details and should be spot-checked before submission.")

# =============================================================================
# order, numbering, save
# =============================================================================
order = [S[0], S[1], S[2], N_new, N_base, S[4], S[8], S[7], S[9], S[10],
         N_two, S[5], S[6], S[3], S[11], N_cmd, S[14], S[15], S[12], S[13],
         P_abl, P_sweep, N_real, N_curves, P_neural, S[16], N_pit, S[17], S[18], P_concl, S[19], S[20]]
assert len(order) == len(set(id(x) for x in order)) == len(prs.slides._sldIdLst), (len(order), len(prs.slides._sldIdLst))
reorder(prs, order)
update_slide_numbers(prs)
prs.save(DST)
print('saved', DST, len(prs.slides._sldIdLst), 'slides')
