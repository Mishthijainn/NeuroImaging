"""Helpers that clone the look of docs/NeuroImaging_CAFN_Presentation.pptx
(Cambria titles, Calibri body, white rounded cards with a soft shadow, theme
colours tx1/tx2/bg1/bg2/accent1..6).  All geometry is in EMU on a 12192000 x
6858000 slide.  Content area: x 548640..11643055, y 1325880..6217920.
"""
import copy, io
from xml.sax.saxutils import escape
from pptx.oxml import parse_xml
from pptx.util import Emu

NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
      'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')
L, W, TOP, BOT = 548640, 11094415, 1325880, 6217920
SHADOW = ('<a:effectLst><a:outerShdw sx="100000" sy="100000" kx="0" ky="0" algn="bl" rotWithShape="0" '
          'blurRad="101600" dist="25400" dir="5400000"><a:srgbClr val="12232B"><a:alpha val="12000"/>'
          '</a:srgbClr></a:outerShdw></a:effectLst>')


def clr(spec):
    """'tx1' | ('accent3', 20) -> pale tint (lumMod 20 / lumOff 80)."""
    if isinstance(spec, tuple):
        name, pct = spec
        return (f'<a:schemeClr val="{name}"><a:lumMod val="{pct*1000}"/>'
                f'<a:lumOff val="{(100-pct)*1000}"/></a:schemeClr>')
    return f'<a:schemeClr val="{spec}"/>'


def nid(slide):
    ids = [int(e.get('id')) for e in slide.shapes._spTree.iter('{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')]
    return max(ids + [1]) + 1


def run(text, sz=14, b=False, i=False, color='tx2', font='Calibri', sub=False):
    attrs = f' sz="{int(sz*100)}"' + (' b="1"' if b else '') + (' i="1"' if i else '') + (' baseline="-25000"' if sub else '')
    return (f'<a:r><a:rPr lang="en-US"{attrs} dirty="0"><a:solidFill>{clr(color)}</a:solidFill>'
            f'<a:latin typeface="{font}" pitchFamily="34" charset="0"/></a:rPr><a:t>{escape(text)}</a:t></a:r>')


def para(runs, algn=None, bullet=False, spc_aft=None, sz_end=14):
    if isinstance(runs, str):
        runs = [runs]
    a = f' algn="{algn}"' if algn else ''
    aft = f'<a:spcAft><a:spcPts val="{int(spc_aft*100)}"/></a:spcAft>' if spc_aft else ''
    if bullet:
        ppr = (f'<a:pPr marL="285750" indent="-285750"{a}>{aft}<a:buSzPct val="100000"/>'
               f'<a:buChar char="&#x2022;"/></a:pPr>')
    else:
        ppr = f'<a:pPr{a} indent="0" marL="0">{aft}<a:buNone/></a:pPr>'
    return f'<a:p>{ppr}{"".join(runs)}<a:endParaRPr lang="en-US" sz="{int(sz_end*100)}" dirty="0"/></a:p>'


def _add(slide, xml):
    el = parse_xml(xml)
    slide.shapes._spTree.append(el)
    return el


def textbox(slide, name, x, y, w, h, paras, anchor='t'):
    if isinstance(paras, str):
        paras = [paras]
    return _add(slide, f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{nid(slide)}" name="{escape(name)}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{int(x)}" y="{int(y)}"/><a:ext cx="{int(w)}" cy="{int(h)}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/><a:ln><a:noFill/></a:ln></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="{anchor}"/><a:lstStyle/>'
        f'{"".join(paras)}</p:txBody></p:sp>')


def shape(slide, name, x, y, w, h, prst='roundRect', adj=4000, fill='bg1', line=None, line_w=19050,
          shadow=True, dash=None, paras=None, anchor='ctr', ins=91440):
    av = f'<a:avLst><a:gd name="adj" fmla="val {adj}"/></a:avLst>' if prst == 'roundRect' else '<a:avLst/>'
    f = f'<a:solidFill>{clr(fill)}</a:solidFill>' if fill else '<a:noFill/>'
    d = f'<a:prstDash val="{dash}"/>' if dash else ''
    ln = f'<a:ln w="{line_w}"><a:solidFill>{clr(line)}</a:solidFill>{d}</a:ln>' if line else '<a:ln><a:noFill/></a:ln>'
    eff = SHADOW if shadow else ''
    body = ''
    if paras:
        body = (f'<p:txBody><a:bodyPr wrap="square" lIns="{ins}" tIns="{ins//2}" rIns="{ins}" bIns="{ins//2}" rtlCol="0" anchor="{anchor}"/>'
                f'<a:lstStyle/>{"".join(paras)}</p:txBody>')
    return _add(slide, f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{nid(slide)}" name="{escape(name)}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{int(x)}" y="{int(y)}"/><a:ext cx="{int(w)}" cy="{int(h)}"/></a:xfrm>'
        f'<a:prstGeom prst="{prst}">{av}</a:prstGeom>{f}{ln}{eff}</p:spPr>{body}</p:sp>')


def card(slide, name, x, y, w, h, adj=4000, fill='bg1'):
    return shape(slide, name, x, y, w, h, 'roundRect', adj, fill)


def picture(slide, name, blob, x, y, w, h, descr=''):
    pic = slide.shapes.add_picture(io.BytesIO(blob), Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    pic.name = name
    pic._element.nvPicPr.cNvPr.set('descr', descr or name)
    return pic


def badge(slide, name, x, y, d, color, icon_blob=None):
    shape(slide, name + ' badge', x, y, d, d, 'ellipse', fill=color, shadow=False)
    if icon_blob:
        k = int(d * 0.5)
        picture(slide, name + ' icon', icon_blob, x + (d - k) / 2, y + (d - k) / 2, k, k)


def arrow(slide, name, x, y, w, h, kind='rightArrow', fill='tx2'):
    return shape(slide, name, x, y, w, h, kind, fill=fill, shadow=False)


def remove(slide, el):
    """Remove a shape/element and drop the image relationship it used."""
    el = getattr(el, '_element', el)
    rids = el.xpath('.//a:blip/@r:embed')
    el.getparent().remove(el)
    for rid in rids:
        try:
            slide.part.drop_rel(rid)
        except Exception:
            pass


def find(slide, prefix=None, name=None, nth=0):
    out = []
    for sh in slide.shapes:
        if name is not None and sh.name == name:
            out.append(sh)
        elif prefix is not None and sh.has_text_frame and sh.text_frame.text.startswith(prefix):
            out.append(sh)
    if len(out) <= nth:
        raise KeyError((prefix, name))
    return out[nth]


def set_text(sh, text):
    """Replace text of a one-paragraph shape keeping the first run's formatting."""
    tf = sh.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p = tf.paragraphs[0]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_paras(sh, paras):
    txb = sh._element.txBody
    for p in txb.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}p'):
        txb.remove(p)
    for p in paras:
        txb.append(parse_xml(p.replace('<a:p>', f'<a:p {NS}>', 1)))


def geom(sh, x=None, y=None, w=None, h=None):
    if x is not None: sh.left = Emu(int(x))
    if y is not None: sh.top = Emu(int(y))
    if w is not None: sh.width = Emu(int(w))
    if h is not None: sh.height = Emu(int(h))


def flow_row(slide, tag, x, y, w, h, labels, color, gap=228600, sz=12, bw=None):
    """Row of rounded boxes joined by right arrows (tinted fill, coloured outline)."""
    n = len(labels)
    colors = color if isinstance(color, list) else [color] * n
    bw = bw or (w - (n - 1) * gap) / n
    for k, lab in enumerate(labels):
        bx = x + k * (bw + gap)
        color = colors[k]
        shape(slide, f'{tag} {k+1}', bx, y, bw, h, 'roundRect', 8000, (color, 18), line=color, line_w=19050,
              shadow=False, paras=[para([run(lab, sz, color='tx1')], algn='ctr', sz_end=sz)])
        if k < n - 1:
            aw = gap * 0.7
            arrow(slide, f'{tag} arrow {k+1}', bx + bw + (gap - aw) / 2, y + h / 2 - 70000, aw, 140000)


def table(slide, name, x, y, col_w, rows, row_h, hdr_h=None, sz=12, hdr_sz=12, first_bold=True,
          algns=None, fills=None):
    """rows[0] is the header. fills: dict row_index -> colour spec for highlighted rows."""
    algns = algns or ['l'] * len(col_w)
    fills = fills or {}
    hdr_h = hdr_h or row_h
    grid = ''.join(f'<a:gridCol w="{int(c)}"/>' for c in col_w)
    trs = ''
    for ri, row in enumerate(rows):
        h = hdr_h if ri == 0 else (row_h[ri - 1] if isinstance(row_h, (list, tuple)) else row_h)
        tcs = ''
        for ci, cell in enumerate(row):
            if ri == 0:
                r_ = run(cell, hdr_sz, b=True, color='bg1'); fill = 'tx2'
            else:
                r_ = run(cell, sz, b=(first_bold and ci == 0), color='tx1')
                fill = fills.get(ri, 'bg1' if ri % 2 else 'bg2')
            ln = ''.join(f'<a:{s} w="0" cap="flat" cmpd="sng" algn="ctr"><a:noFill/></a:{s}>' for s in ('lnL', 'lnR', 'lnT', 'lnB'))
            tcs += (f'<a:tc><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:pPr algn="{algns[ci]}" indent="0" marL="0"><a:buNone/></a:pPr>{r_}'
                    f'<a:endParaRPr lang="en-US" sz="{int(sz*100)}" dirty="0"/></a:p></a:txBody>'
                    f'<a:tcPr marL="91440" marR="91440" marT="45720" marB="45720" anchor="ctr">{ln}<a:solidFill>{clr(fill)}</a:solidFill></a:tcPr></a:tc>')
        trs += f'<a:tr h="{int(h)}">{tcs}</a:tr>'
    total_h = hdr_h + (sum(row_h) if isinstance(row_h, (list, tuple)) else row_h * (len(rows) - 1))
    return _add(slide, f'<p:graphicFrame {NS}><p:nvGraphicFramePr><p:cNvPr id="{nid(slide)}" name="{escape(name)}"/>'
        f'<p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>'
        f'<p:xfrm><a:off x="{int(x)}" y="{int(y)}"/><a:ext cx="{int(sum(col_w))}" cy="{int(total_h)}"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table"><a:tbl><a:tblPr/>'
        f'<a:tblGrid>{grid}</a:tblGrid>{trs}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>')


# ---- new slides -------------------------------------------------------------------
class SlideFactory:
    """Creates blank CONTENT slides carrying the standard title + slide-number shapes
    cloned from a template slide of the original deck."""

    def __init__(self, prs, template_slide):
        self.prs = prs
        self.layout = template_slide.slide_layout
        self.title_xml = None
        self.num_xml = None
        for sh in template_slide.shapes:
            if sh.is_placeholder and sh.placeholder_format.type is not None:
                t = str(sh.placeholder_format.type)
                if 'TITLE' in t:
                    self.title_el = copy.deepcopy(sh._element)
                elif 'SLIDE_NUMBER' in t:
                    self.num_el = copy.deepcopy(sh._element)

    def new(self, title, dark=False):
        s = self.prs.slides.add_slide(self.layout)
        for sh in list(s.shapes):
            sh._element.getparent().remove(sh._element)
        t = copy.deepcopy(self.title_el)
        s.shapes._spTree.append(t)
        n = copy.deepcopy(self.num_el)
        s.shapes._spTree.append(n)
        ts = [sh for sh in s.shapes if sh.is_placeholder and 'TITLE' in str(sh.placeholder_format.type)][0]
        set_text(ts, title)
        return s


def move_slide(prs, slide, new_index):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    for el in ids:
        if prs.part.related_part(el.rId) is slide.part:
            lst.remove(el)
            lst.insert(new_index, el)
            return
    raise KeyError


def reorder(prs, slides_in_order):
    lst = prs.slides._sldIdLst
    by_part = {id(prs.part.related_part(el.rId)): el for el in list(lst)}
    for el in list(lst):
        lst.remove(el)
    for s in slides_in_order:
        lst.append(by_part[id(s.part)])


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def set_color(sh, color):
    for r in sh._element.xpath('.//a:rPr/a:solidFill'):
        for c in list(r):
            r.remove(c)
        r.append(parse_xml(f'<a:schemeClr {NS} val="{color}"/>'))


def update_slide_numbers(prs):
    for i, s in enumerate(prs.slides, 1):
        for t in s.shapes._spTree.xpath('.//a:fld[@type="slidenum"]/a:t'):
            t.text = str(i)
